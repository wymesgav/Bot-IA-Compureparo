"""Carga el catalogo de precios de Compureparo desde el catalogo.xlsx en S3.

Wilmar edita el Excel y lo vuelve a subir al bucket; no hace falta reconstruir
la imagen del contenedor para cambiar un precio.

Si S3 falla (bucket sin configurar, sin permisos, archivo corrupto), se usa la
copia incluida en la imagen como respaldo, para que el bot nunca quede mudo.
"""

import io
import logging
import os
import time

from openpyxl import load_workbook

log = logging.getLogger()

CATALOGO_BUCKET = os.environ.get("CATALOGO_BUCKET", "")
CATALOGO_KEY = os.environ.get("CATALOGO_KEY", "catalogo.xlsx")
CATALOGO_LOCAL = os.path.join(os.path.dirname(__file__), "catalogo.xlsx")

# Cuanto tiempo reusar el catalogo ya cargado antes de volver a pedirlo a S3.
# Con esto un cambio de precio se refleja en <=5 min sin leer S3 en cada mensaje.
TTL_SEGUNDOS = int(os.environ.get("CATALOGO_TTL", "300"))

_cache = {"texto": None, "cargado_en": 0.0, "origen": None}

_s3 = None


def _cliente_s3():
    """Crea el cliente S3 solo si de verdad se va a usar."""
    global _s3
    if _s3 is None:
        import boto3

        _s3 = boto3.client("s3")
    return _s3


def obtener_catalogo():
    """Devuelve el catalogo como texto listo para inyectar en el system prompt."""
    ahora = time.time()
    if _cache["texto"] and (ahora - _cache["cargado_en"]) < TTL_SEGUNDOS:
        return _cache["texto"]

    texto, origen = _cargar()
    if texto:
        _cache.update(texto=texto, cargado_en=ahora, origen=origen)
        log.info("Catalogo cargado desde %s (%d caracteres)", origen, len(texto))
        return texto

    # Ultimo recurso: conservar lo que ya teniamos en cache aunque este vencido.
    if _cache["texto"]:
        log.warning("No se pudo recargar el catalogo; se reusa la copia en cache")
        return _cache["texto"]

    log.error("No se pudo cargar el catalogo por ningun medio")
    return ""


def _cargar():
    if CATALOGO_BUCKET:
        try:
            obj = _cliente_s3().get_object(Bucket=CATALOGO_BUCKET, Key=CATALOGO_KEY)
            return _xlsx_a_texto(io.BytesIO(obj["Body"].read())), f"s3://{CATALOGO_BUCKET}/{CATALOGO_KEY}"
        except Exception:
            log.exception("Fallo al leer el catalogo desde S3; se intenta la copia local")

    try:
        with open(CATALOGO_LOCAL, "rb") as f:
            return _xlsx_a_texto(io.BytesIO(f.read())), "copia local de la imagen"
    except Exception:
        log.exception("Fallo al leer la copia local del catalogo")
        return None, None


def _precio(valor):
    """Formatea el precio como lo escribiria Wilmar: $250.000"""
    if valor is None or valor == "":
        return "SIN PRECIO"
    if isinstance(valor, (int, float)):
        return f"${valor:,.0f}".replace(",", ".")
    return str(valor).strip()


def _xlsx_a_texto(binario):
    wb = load_workbook(binario, data_only=True)
    partes = []

    if "Servicios" in wb.sheetnames:
        partes.append("### Catalogo de precios\n")
        categoria_actual = None
        for fila in wb["Servicios"].iter_rows(min_row=2, values_only=True):
            categoria, item, precio, incluye, notas = (list(fila) + [None] * 5)[:5]
            if not item:
                continue
            if categoria != categoria_actual:
                categoria_actual = categoria
                partes.append(f"\n**{categoria}**")

            precio_txt = _precio(precio)
            linea = f"- {item}: {precio_txt}"
            if incluye:
                linea += f"\n  - Incluye: {incluye}"
            if notas:
                linea += f"\n  - Nota: {notas}"
            partes.append(linea)

    if "InfoNegocio" in wb.sheetnames:
        partes.append("\n\n### Informacion del negocio\n")
        for tema, contenido in wb["InfoNegocio"].iter_rows(min_row=2, values_only=True):
            if tema:
                partes.append(f"- **{tema}:** {contenido or ''}")

    if "Diagnosticos" in wb.sheetnames:
        partes.append("\n\n### Guias de diagnostico que puedes dar al cliente\n")
        for necesita, pasos in wb["Diagnosticos"].iter_rows(min_row=2, values_only=True):
            if necesita:
                partes.append(f"- **{necesita}:** {pasos or ''}")

    return "\n".join(partes)
