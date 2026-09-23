"""Genera catalogo.xlsx para el bot de Compureparo.

Fuente de los datos: respuestas rapidas reales del WhatsApp Business de Compureparo
(leidas el 2026-09-22). Las filas marcadas PENDIENTE las completa Wilmar.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ---------------------------------------------------------------- estilos
AZUL = "1F4E79"
GRIS = "F2F2F2"
AMARILLO = "FFF2CC"

h_font = Font(bold=True, color="FFFFFF", size=11)
h_fill = PatternFill("solid", fgColor=AZUL)
h_align = Alignment(vertical="center", horizontal="center", wrap_text=True)

borde = Border(*[Side(style="thin", color="BFBFBF")] * 4)
wrap = Alignment(vertical="top", wrap_text=True)
pendiente_fill = PatternFill("solid", fgColor=AMARILLO)


def hoja(ws, encabezados, filas, anchos):
    ws.append(encabezados)
    for c in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=c)
        celda.font = h_font
        celda.fill = h_fill
        celda.alignment = h_align
        celda.border = borde
    ws.row_dimensions[1].height = 30

    for fila in filas:
        ws.append(fila)

    for r in range(2, ws.max_row + 1):
        es_pendiente = any(
            isinstance(v, str) and "PENDIENTE" in v
            for v in [ws.cell(row=r, column=c).value for c in range(1, len(encabezados) + 1)]
        )
        for c in range(1, len(encabezados) + 1):
            celda = ws.cell(row=r, column=c)
            celda.alignment = wrap
            celda.border = borde
            if es_pendiente:
                celda.fill = pendiente_fill

    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.freeze_panes = "A2"


# ---------------------------------------------------------------- SERVICIOS
ws = wb.active
ws.title = "Servicios"

ENC = ["Categoria", "Item", "Precio (COP)", "Incluye / Condiciones", "Notas para el bot"]

servicios = [
    # --- SSD SATA Windows
    ["SSD SATA (Windows)", "SSD 256 GB", 250000,
     "Instalacion de Windows y Office con licencia permanente. 1 ano de garantia en el SSD. Incluye limpieza interna y cambio de pasta termica.",
     "Advertir que los SSD subieron de precio a nivel mundial."],
    ["SSD SATA (Windows)", "SSD 512 GB", 390000,
     "Instalacion de Windows y Office con licencia permanente. 1 ano de garantia en el SSD. Incluye limpieza interna y cambio de pasta termica.",
     ""],
    ["SSD SATA (Windows)", "SSD 1 TB", 650000,
     "Instalacion de Windows y Office con licencia permanente. 1 ano de garantia en el SSD. Incluye limpieza interna y cambio de pasta termica.",
     ""],

    # --- SSD SATA Mac
    ["SSD SATA (Mac)", "SSD 240 GB", 230000,
     "Incluye instalacion del sistema operativo MacOS. 1 ano de garantia.", ""],
    ["SSD SATA (Mac)", "SSD 512 GB", 330000,
     "Incluye instalacion del sistema operativo MacOS. 1 ano de garantia.", ""],
    ["SSD SATA (Mac)", "SSD 960 GB", 420000,
     "Incluye instalacion del sistema operativo MacOS. 1 ano de garantia.", ""],

    # --- SSD M.2 NVMe
    ["SSD M.2 NVMe", "SSD 240 GB NVMe PCIe M.2 Gen3 x4", 230000,
     "Incluye instalacion de Windows y Office. 1 ano de garantia en el SSD.", ""],
    ["SSD M.2 NVMe", "SSD 512 GB M.2 PCIe Gen3 x4", 310000,
     "Incluye instalacion de Windows y Office. 1 ano de garantia en el SSD.", ""],
    ["SSD M.2 NVMe", "SSD 1 TB M.2 PCIe Gen3 x4", 420000,
     "Incluye instalacion de Windows y Office. 1 ano de garantia en el SSD.", ""],

    # --- Office
    ["Software", "Office 2016 Pro - licencia permanente por codigo", 30000,
     "Solo Windows. Se instala de forma remota.", ""],
    ["Software", "Office 2019 o 2021 Pro - licencia original de volumen", 110000,
     "Licencia permanente y original con Microsoft (volumen = al por mayor, por eso es mas economica). Se instala de forma remota.",
     ""],

    # --- Pantallas
    ["Pantallas", 'Pantalla 15,6" 30 pines Matte', 340000,
     "Instalacion incluida. Matte = antireflejo.",
     "Precio aproximado: la resolucion exacta se confirma con el numero de parte al quitar la pantalla."],
    ["Pantallas", 'Pantalla 15,6" glossy 1366x768', 260000,
     "Instalacion incluida.",
     "Precio aproximado: la resolucion exacta se confirma con el numero de parte al quitar la pantalla."],

    # --- Domicilio
    ["Servicio a domicilio", "Revision a domicilio (solo si NO se repara)", 20000,
     "Se revisa en sitio y se diagnostica la falla. Si se da solucion en la visita, el cliente solo paga la reparacion. Si no acepta la reparacion o no se presta para ello, paga solo este valor.",
     "Si no es posible resolver en sitio, se recoge el equipo para un diagnostico mas completo."],

    # --- Sin precio documentado
    ["Reparaciones", "Cambio de teclado con remaches", "PENDIENTE",
     "Requiere recoger el equipo. Entrega el mismo dia o el dia siguiente (se debe poner a calor con herramienta).",
     "No hay precio documentado: el bot debe escalar a Wilmar."],
    ["Mantenimiento", "Mantenimiento / limpieza", "PENDIENTE", "", "Sin precio confirmado: el bot escala a Wilmar."],
    ["Redes", "Instalacion de red / punto de red", "PENDIENTE", "", "Sin precio confirmado: el bot escala a Wilmar."],
    ["Camaras", "Instalacion de camaras de seguridad", "PENDIENTE", "", "Sin precio confirmado: el bot escala a Wilmar."],
    ["Memoria", "Ampliacion de memoria RAM", "PENDIENTE", "", "Sin precio confirmado: el bot escala a Wilmar."],
    ["Software", "Formateo / reinstalacion de sistema operativo", "PENDIENTE", "", "Sin precio confirmado: el bot escala a Wilmar."],
]

hoja(ws, ENC, servicios, [22, 38, 15, 55, 45])
for r in range(2, ws.max_row + 1):
    celda = ws.cell(row=r, column=3)
    if isinstance(celda.value, int):
        celda.number_format = '"$"#,##0'

# ---------------------------------------------------------------- INFO NEGOCIO
ws2 = wb.create_sheet("InfoNegocio")
info = [
    ["Ubicacion", "Estamos en Itagui. El servicio es a domicilio, o con recogida del equipo para entrega posterior."],
    ["Como funciona el domicilio",
     "Primero se revisa en sitio y se diagnostica la falla. Si se da solucion en la visita, solo se paga la reparacion. Si no acepta la reparacion o no se presta para ello, solo paga $20.000. Si no es posible resolver en sitio, se recoge el equipo para diagnostico completo."],
    ["Soporte remoto",
     "El cliente instala TeamViewer o AnyDesk y da acceso. Si no sabe instalarlo, se le envia un video."],
    ["Quien atiende", "Wilmar Mesa - tecnico e Ingeniero de Telecomunicaciones."],
    ["Despedida habitual", "Recuerda que hablaste con Wilmar, cualquier inquietud con gusto te ayudare. Feliz dia."],
    ["Datos de pago",
     "NO los entrega el bot. Si el cliente pregunta como pagar, el bot escala a Wilmar."],
]
hoja(ws2, ["Tema", "Contenido"], info, [28, 100])

# ---------------------------------------------------------------- DIAGNOSTICOS
ws3 = wb.create_sheet("Diagnosticos")
diag = [
    ["Saber si el disco es SSD o HDD (Windows 11)",
     "1. Clic derecho en la barra de tareas. 2. Selecciona 'Administrador de tareas'. 3. Pestana 'Rendimiento'. 4. Selecciona 'Disco' (puede aparecer como Disco 0, HDD, SSD). 5. Envia una foto."],
    ["Ver la memoria RAM y los slots",
     "1. Clic derecho en la barra de tareas. 2. Administrador de tareas. 3. Pestana Rendimiento. 4. Clic en Memoria o RAM. 5. Envia una foto."],
    ["Ver la board / informacion del equipo",
     "1. Presiona Win + R para abrir Ejecutar. 2. Escribe msinfo32 y presiona Enter. 3. Envia una foto."],
]
hoja(ws3, ["Que necesita saber el cliente", "Pasos que le indica el bot"], diag, [42, 100])

ruta = r"C:\Users\Will\OneDrive - INSTITUTO TECNOLOGICO METROPOLITANO - ITM\Desktop\Claude\Proyectos\Perfil\Bot IA\catalogo.xlsx"
wb.save(ruta)
print("OK ->", ruta)
